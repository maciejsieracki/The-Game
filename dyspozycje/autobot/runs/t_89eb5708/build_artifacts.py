from __future__ import annotations
from collections import Counter
from hashlib import sha256
from pathlib import Path
from datetime import datetime
import json, os
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path.cwd()
TASK = "t_89eb5708"
RUN_AT = os.environ.get("RUN_AT") or datetime.now().astimezone().isoformat()
OUT = ROOT / "panele-sterowania" / "cyw-macierz"
RUN = ROOT / "dyspozycje" / "autobot" / "runs" / TASK
OUT.mkdir(parents=True, exist_ok=True)
RUN.mkdir(parents=True, exist_ok=True)


def load(rel: str):
    return json.loads((ROOT / rel).read_text(encoding="utf-8"))

matrix = load("gra/data/civ-matrix.json")
civs_data = load("gra/data/civs.json")
ai_data = load("gra/data/civ-ai.json")
civ_params_data = load("gra/data/civ-params.json")
dip_data = load("gra/data/diplomacy.json")
all_rows = matrix["cywilizacje"]
target_rows = [r for r in all_rows if r["typCywilizacji"] != "grecy"]
param_defs = matrix["paramDefs"]
defaults = matrix["defaults"]
param_ids = list(param_defs)
target_ids = [r["typCywilizacji"] for r in target_rows]
assert len(target_rows) == 14
assert len(param_ids) == 113
assert len(set(target_ids)) == 14 and "grecy" not in target_ids
assert all(set(r["params"]) == set(param_ids) for r in target_rows)

legacy_docs = {
    "rzymianie": "docs/encyklopedia/cywilizacje/rzymianie.md",
    "chinczycy": "docs/encyklopedia/cywilizacje/chinczycy.md",
    "inkowie": "docs/encyklopedia/cywilizacje/inkowie.md",
    "zulusi": "docs/encyklopedia/cywilizacje/zulusi.md",
    "egipt": "docs/encyklopedia/cywilizacje/egipt.md",
    "sumer": "docs/encyklopedia/cywilizacje/sumer.md",
    "celtowie": "docs/encyklopedia/cywilizacje/celtowie.md",
    "germanie": "docs/encyklopedia/cywilizacje/germanie.md",
}
reserve_ids = {"harappa", "hetyci", "slowianie", "babilonia", "asyria", "fenicjanie"}
assert set(target_ids) == set(legacy_docs) | reserve_ids

AI = {
    "ai_agresywnosc": ("REAL_GAMEPLAY", "SKALA_1_10", "gra/src/game/civ-ai-data.ts:38-53,156-178; gra/src/game/diplomacy.ts:1261-1266; gra/src/main.ts:19784-19790,33910-33915; gra/src/game/ai.ts:4974-5036,5184-5187", "wyżej = bardziej agresywny AI; Hard +1, Easy -1", "easy=max(1,normal-1); normal=normal; hard=min(10,normal+1); /10", "Matrix row shadows civ-ai.json."),
    "ai_ekspansywnosc": ("REAL_GAMEPLAY", "SKALA_1_10", "gra/src/game/civ-ai-data.ts:74-102; gra/src/main.ts:33264-33296; gra/src/game/ai.ts:506-513,1097-1117,2857-2868,4257-4264", "wyżej = większa ekspansja; Hard +1, Easy -1", "easy=max(1,normal-1); normal=normal; hard=min(10,normal+1)", "Used by city founding and cluster consolidation."),
    "ai_priorytet_militarny": ("REAL_GAMEPLAY", "SKALA_1_10", "gra/src/game/civ-ai-data.ts:74-102; gra/src/main.ts:33264-33296; gra/src/game/ai-production-priorities.ts:6-27; gra/src/game/ai.ts:1672-1676", "wyżej = większy score wojskowy; Hard +1, Easy -1", "easy=max(1,normal-1); normal=normal; hard=min(10,normal+1); delta=(value-5)*15", "AI production score."),
    "ai_priorytet_ekonomia": ("REAL_GAMEPLAY", "SKALA_1_10", "gra/src/game/civ-ai-data.ts:74-102; gra/src/main.ts:33264-33296; gra/src/game/ai-production-priorities.ts:6-27; gra/src/game/ai.ts:1672-1676", "wyżej = większy score ekonomiczny; Hard +1, Easy -1", "easy=max(1,normal-1); normal=normal; hard=min(10,normal+1); delta=(value-5)*15", "AI production score."),
    "ai_priorytet_nauka": ("REAL_GAMEPLAY", "SKALA_1_10", "gra/src/game/civ-ai-data.ts:74-102; gra/src/main.ts:33264-33296; gra/src/game/ai-production-priorities.ts:6-27; gra/src/game/ai.ts:1672-1676", "wyżej = większy score nauki; Hard +1, Easy -1", "easy=max(1,normal-1); normal=normal; hard=min(10,normal+1); delta=(value-5)*15", "AI production score."),
    "ai_tolerancja_ryzyka": ("REAL_GAMEPLAY", "SKALA_1_10", "gra/src/game/civ-ai-data.ts:74-102; gra/src/main.ts:33924-33927; gra/src/game/ai.ts:4974-5036,5184-5187", "wyżej = większa tolerancja ryzyka wojny; Hard +1, Easy -1", "easy=max(1,normal-1); normal=normal; hard=min(10,normal+1)", "Diplomacy bias."),
    "ai_sklonnosc_podboju": ("REAL_GAMEPLAY", "SKALA_1_10", "gra/src/game/civ-ai-data.ts:74-102; gra/src/main.ts:33264-33296,33924-33927; gra/src/game/ai.ts:4974-5036,5184-5187", "wyżej = większa skłonność do podboju; Hard +1, Easy -1", "easy=max(1,normal-1); normal=normal; hard=min(10,normal+1)", "Founding/conquest and diplomacy bias."),
    "ai_profil_obronna": ("REAL_GAMEPLAY", "FLAG_IDENTITY", "gra/src/game/civ-ai-data.ts:56-71,119-123; gra/src/main.ts:9169-9173", "1 = defensive-copy map profile; no difficulty delta", "flag raw 0/1 at Easy/Normal/Hard", "Owner-role identity flag."),
}
UI = {
    "dip_otwartosc_handel": "gra/src/game/diplomacy-display.ts:44-49,116-122",
    "dip_sklonnosc_sojusze": "gra/src/game/diplomacy-display.ts:54-59,116-122",
    "dip_lojalnosc": "gra/src/game/diplomacy-display.ts:64-69,116-122",
    "dip_prog_wojny": "gra/src/game/diplomacy-display.ts:73-78,116-122",
    "dip_pamietliwosc": "gra/src/game/diplomacy-display.ts:82-87,116-122",
}
SPECIAL = {
    "lud_wzrost_proc": ("REAL_GAMEPLAY", "GROWTH_DIFFICULTY_MULT", "gra/src/game/population-growth-v85.ts:84-90,230-232", "wyżej = szybszy wzrost i silniejszy AI; Hard ×1.5, Easy ×0.5", "easy=raw*0.5; normal=raw; hard=raw*1.5; contribution rounded to percent", "Only matrix field explicitly difficulty-scaled by growth consumer."),
    "dip_handlowosc_archetyp": ("REAL_GAMEPLAY", "CURRENT_NEUTRAL_MUL_ABS", "gra/src/game/civ-ai-data.ts:184-198; gra/src/game/diplomacy.ts:1267-1270; gra/src/main.ts:19553-19556,33555-33559,33916-33919", "direction requires decision; no arbitrary Easy/Hard", "current easy=raw; normal=raw; hard=raw because formula=mul_abs", "Real consumer; helper scales only formula=skala."),
    "dip_nastawienie_bazowe": ("REAL_GAMEPLAY", "CURRENT_NEUTRAL_CALLER_OMITS_DIFFICULTY", "gra/src/game/civ-ai-data.ts:205-217; gra/src/game/diplomacy.ts:1376-1380", "direction requires decision; no arbitrary Easy/Hard", "current easy=raw; normal=raw; hard=raw; caller uses default normal", "Function accepts difficulty but caller omits it."),
}


def meta(pid: str) -> dict[str, str]:
    if pid in AI:
        s, p, c, d, f, n = AI[pid]
    elif pid in SPECIAL:
        s, p, c, d, f, n = SPECIAL[pid]
    elif pid in UI:
        return {"status": "UI_ONLY", "policy": "DISPLAY_RAW", "calls": UI[pid], "direction": "etykieta UI; brak wpływu na wynik gry", "formula": "display reads raw value; no gameplay Easy/Normal/Hard effect", "note": "Display-only character tag."}
    else:
        non_default = any(r["params"].get(pid) != defaults.get(pid) for r in target_rows)
        return {"status": "DEAD_UNWIRED" if non_default else "UNWIRED", "policy": "NOT_ACTIVE", "calls": "gra/src/game/civ-matrix.ts:41-60,71-101 (loader/accessor only); no gameplay call-site found", "direction": "not applicable until a consumer is wired and reviewed", "formula": "not computed by current consumer; retain raw Normal data", "note": "paramDefs.modul is not proof of a live consumer."}
    return {"status": s, "policy": p, "calls": c, "direction": d, "formula": f, "note": n}


def effective(pid: str, raw: Any, m: dict[str, str]):
    if not isinstance(raw, (int, float)):
        return "", "", ""
    if m["status"] != "REAL_GAMEPLAY":
        return "", raw, ""
    if m["policy"] == "SKALA_1_10":
        return max(1, min(10, raw - 1)), raw, max(1, min(10, raw + 1))
    if m["policy"] == "GROWTH_DIFFICULTY_MULT":
        return raw * 0.5, raw, raw * 1.5
    return raw, raw, raw


def clean(v: Any):
    return json.dumps(v, ensure_ascii=False, sort_keys=True) if isinstance(v, (dict, list)) else v

legacy = {r["Cywilizacja"]: r for r in ai_data["cywilizacje"]}
params = {r["Cywilizacja"]: r for r in civ_params_data["cywilizacje"]}
dips = {r["Cywilizacja"]: r for r in dip_data.get("perNacja", [])}
civs = {r["Cywilizacja"]: r for r in civs_data["cywilizacje"]}

def value_path(cid: str, pid: str):
    return f"gra/data/civ-matrix.json#/cywilizacje[typCywilizacji={cid}]/params/{pid}"

def history(cid: str):
    if cid in legacy_docs:
        return ("LEGACY_DOC_PRESENT", legacy_docs[cid], "current-data-high/historical-text-present", "Encyclopedia entry; matrix value remains current JSON.")
    return ("RESERVE_DRAFT_UNRESOLVED", "docs/decyzje/D-cyw-roster-6-REZERWA.md; Civ-CYWILIZACJE/draft/roster-6-REZERWA.json", "current-data-high/historical-acceptance-open", "Reserve evidence exists; accepted promotion is not claimed.")

matrix_rows = []
coverage = []
status_counts = Counter()
nondefault = 0
for civ in target_rows:
    cid = civ["typCywilizacji"]
    hs, hc, hconf, hnote = history(cid)
    for pid, spec in param_defs.items():
        raw, default = civ["params"][pid], defaults.get(pid, 0)
        m = meta(pid)
        e, n, h = effective(pid, raw, m)
        nd = raw != default
        nondefault += nd
        status_counts[m["status"]] += 1
        prov = "CURRENT_RUNTIME_DATA" if nd else "CURRENT_RUNTIME_DEFAULT"
        common = {
            "civ_id": cid, "civilization": civ["Cywilizacja"], "tier": civ["tier"], "param_id": pid,
            "domain": spec.get("domena", ""), "unit": spec.get("jednostka", ""), "declared_module": spec.get("modul", ""), "formula": spec.get("formula", ""),
            "raw_current": raw, "default": default, "non_default": nd, "value_status": "CURRENT_RUNTIME_DATA", "default_relation": "NON_DEFAULT_OVERRIDE" if nd else "DEFAULT_VALUE",
            "provenance": prov, "value_source": value_path(cid, pid), "historical_status": hs, "historical_source": hc, "consumer_status": m["status"],
            "easy_effective": e, "normal_effective": n, "hard_effective": h, "difficulty_policy": m["policy"], "effective_formula": m["formula"], "direction": m["direction"], "live_call_site": m["calls"], "copied_from_greece": False,
        }
        coverage.append(common)
        matrix_rows.append([cid, civ["Cywilizacja"], civ["tier"], pid, spec.get("domena", ""), spec.get("jednostka", ""), spec.get("modul", ""), spec.get("formula", ""), raw, default, raw-default if isinstance(raw,(int,float)) and isinstance(default,(int,float)) else "", "CURRENT_RUNTIME_DATA", "NON_DEFAULT_OVERRIDE" if nd else "DEFAULT_VALUE", prov, value_path(cid,pid), hc, hconf, m["status"], e, n, h, m["policy"], m["formula"], m["direction"], m["calls"], "copied_from_greece=false; direct target-row lookup; "+m["note"]])
assert len(coverage) == 14 * 113 == 1582
assert all(not x["copied_from_greece"] for x in coverage)

profile_rows = []
ai_map = {
    "ai_agresywnosc": "agresywnosc", "ai_ekspansywnosc": "ekspansywnosc", "ai_priorytet_militarny": "priorytetMilitarny", "ai_priorytet_ekonomia": "priorytetEkonomia", "ai_priorytet_nauka": "priorytetNauka", "ai_tolerancja_ryzyka": "tolerancjaRyzyka", "ai_sklonnosc_podboju": "sklonnoscDoPodboju", "ai_profil_obronna": "profilMapy",
}
for civ in target_rows:
    cid, cname = civ["typCywilizacji"], civ["Cywilizacja"]
    for pid, old_key in ai_map.items():
        raw = civ["params"][pid]; m = meta(pid); e,n,h = effective(pid, raw, m); old = legacy.get(cname, {}).get(old_key, "")
        profile_rows.append([cid,cname,"AI",pid,raw,old,param_defs[pid].get("jednostka",""),m["status"],e,n,h,m["policy"],value_path(cid,pid),"gra/data/civ-ai.json",m["calls"],"matrix overrides legacy adapter"])
    for field in ["preferowaneBudynki","preferowaneJednostki","modWzrostu","modEkonomii","uwagi"]:
        val=params.get(cname,{}).get(field,"")
        profile_rows.append([cid,cname,"CIV_PARAMS",field,val,val,"","DEAD_UNWIRED","",val,"","NOT_ACTIVE",f"gra/data/civ-params.json#/cywilizacje[Cywilizacja={cname}]/{field}","gra/src/game/civ-ai-data.ts:115-117","gra/src/game/civ-ai-data.ts:115-117 only; no downstream consumer found","reference/adapter layer"])
    for field,pid in [("sklonnoscSojusze","dip_sklonnosc_sojusze"),("lojalnosc","dip_lojalnosc"),("progWojny","dip_prog_wojny"),("pamietliwosc","dip_pamietliwosc"),("otwartoscHandel","dip_otwartosc_handel"),("nastawienieBazowe","dip_nastawienie_bazowe")]:
        val=dips.get(cname,{}).get(field,"")
        profile_rows.append([cid,cname,"DIPLOMACY_PERNACJA",field,val,val,"1_10" if field!="nastawienieBazowe" else "absolut","DEAD_UNWIRED","",val,"","FALLBACK_ONLY",f"gra/data/diplomacy.json#/perNacja[Cywilizacja={cname}]/{field}","gra/src/game/civ-ai-data.ts:143-150; gra/src/game/diplomacy-display.ts:116-122","matrix row shadows fallback for all 14 target rows","legacy/fallback provenance"])

bonus_rows=[]
bonus_status=Counter()
for civ in target_rows:
    cid,cname=civ["typCywilizacji"],civ["Cywilizacja"]; cr=civs[cname]
    bonus_rows.append([cid,cname,"civs.json field","mnoznikHandelPieniadz","handel->pieniadz",cr.get("mnoznikHandelPieniadz",""),"RUNTIME_CONSUMER","gra/data/civs.json","gra/src/game/economy.ts:762-817; gra/src/game/turn-economy.ts:292-295","Normal raw; easy +0.5; hard -0.5","Not matrix meta_mnoznik_waluta"]); bonus_status["RUNTIME_CONSUMER"]+=1
    for i,b in enumerate(cr.get("bonusy",[]),1):
        typ,real=b.get("typ",""),b.get("realizuje","")
        if typ in {"bonus_walka","bonus_obrona","koszt_redukcja"} and real in {"walka","miasto"}: st,call="RUNTIME_CONSUMER","gra/src/game/civ-bonuses.ts:130-176,232-267"
        elif typ in {"bonus_zloto","bonus_nauka"} and real=="ekonomia": st,call="RUNTIME_CONSUMER","gra/src/game/economy.ts:820-855"
        elif typ in {"mnoznik_manpower_max","bonus_pobor_regen"}: st,call="RUNTIME_CONSUMER","gra/src/game/manpower.ts:370-405"
        else: st,call="UI_ONLY","gra/data/civs.json; no matching numeric consumer established"
        bonus_status[st]+=1
        bonus_rows.append([cid,cname,"civs.json bonusy",f"bonusy[{i}]",typ,clean(b.get("cel","")),clean(b.get("wartosc","")),st,"gra/data/civs.json",call,"raw bonus; no matrix substitution",b.get("opis","")])

policy_rows=[]
for pid in param_ids:
    m=meta(pid); raw=target_rows[0]["params"][pid]; e,n,h=effective(pid,raw,m)
    policy_rows.append([pid,param_defs[pid].get("domena",""),m["status"],m["policy"],raw,e,n,h,m["formula"],m["direction"],m["calls"],m["note"]])

allocation_rows=[
 ["grecy","major-ai","easy",50,50,60,20,20,100,"gra/src/game/civ-ai-allocation.ts:34-55","Pilot profile is explicitly difficulty-neutral."],
 ["grecy","major-ai","normal",50,50,60,20,20,100,"gra/src/game/civ-ai-allocation.ts:34-55","Pilot profile is explicitly difficulty-neutral."],
 ["grecy","major-ai","hard",50,50,60,20,20,100,"gra/src/game/civ-ai-allocation.ts:34-55","Pilot profile is explicitly difficulty-neutral."],
 ["*","city-state","easy/normal/hard","n/a","n/a","n/a","n/a","n/a",100,"gra/src/game/civ-ai-allocation.ts:75-87","AI-wide envelope."],
 ["*","defensive-copy","easy/normal/hard","n/a","n/a","n/a","n/a","n/a",100,"gra/src/game/civ-ai-allocation.ts:75-87","AI-wide envelope."],
 ["*","player/hotseat","easy/normal/hard","n/a","n/a","n/a","n/a","n/a",33,"gra/src/game/civ-ai-allocation.ts:26-32,75-87","Human automation envelope."],
 ["14 target civs","major-ai","easy/normal/hard","not defined","not defined","not defined","not defined","not defined","AI-wide fallback 100","gra/src/game/civ-ai-allocation.ts:61-67,86-87","No Greece rows copied; separate approval needed for per-civ profiles."],
]
history_rows=[]
for civ in target_rows:
    cid,cname=civ["typCywilizacji"],civ["Cywilizacja"]; hs,hc,hconf,hnote=history(cid)
    history_rows.append([cid,cname,civ["tier"],hs,hc,hnote,hconf,f"gra/data/civ-matrix.json#/cywilizacje[typCywilizacji={cid}]","CURRENT_RUNTIME_DATA","No accepted historical promotion claimed." if cid in reserve_ids else "Current matrix is runtime evidence; encyclopedia is text evidence."])

unresolved=[
 ["HIST-14-RESERVE","Six reserve civilizations","BLOCKED","Six reserve IDs have draft evidence, but this artifact does not promote them to accepted historical spec.","Owner/Final Control confirms accepted status or keeps source-gap.","docs/decyzje/D-cyw-roster-6-REZERWA.md; Civ-CYWILIZACJE/draft/roster-6-REZERWA.json"],
 ["WIRE-113","Unwired matrix fields","DEAD_UNWIRED","All 113 rows load, but only traced consumer fields are live; declarative modul is not enough.","Break into implementation topics; do not wire in this evidence task.","gra/src/game/civ-matrix.ts:41-101; Konsumenci_i_braki"],
 ["DIFF-TRADE","dip_handlowosc_archetyp difficulty","DECISION_REQUIRED","Real consumer is current-neutral because mul_abs is not scaled by helper.","Choose neutral-by-design or approve Easy/Normal/Hard scaling before code.","gra/src/game/civ-ai-data.ts:184-198; gra/src/game/civ-matrix.ts:63-79"],
 ["DIFF-TRUST","dip_nastawienie_bazowe difficulty","DECISION_REQUIRED","Caller omits difficulty and therefore uses Normal.","Choose neutral-by-design or approve scaling and pass difficulty before code.","gra/src/game/civ-ai-data.ts:205-217; gra/src/game/diplomacy.ts:1376-1380"],
 ["ALLOC-14","Per-civilization allocation","PROPOSAL","Only Grecy have a civilization-specific allocation profile; other 14 use fallback.","Do not copy Greece; create separate approved allocation topic if needed.","gra/src/game/civ-ai-allocation.ts:34-87"],
 ["SOURCE-XLSX","Canonical historical workbook","UNRESOLVED","Base has JSON export/schema/exporter but no tracked original 11-sheet XLSX.","Attach authoritative workbook or immutable citation if archive provenance is required.","gra/data/civ-matrix.json._meta.source"],
]

# Source manifest and exact source hits.
source_paths=['gra/data/civ-matrix.json','gra/data/civs.json','gra/data/civ-ai.json','gra/data/civ-params.json','gra/data/diplomacy.json','gra/data/units.json','panele-sterowania/cyw_macierz_schema.py','panele-sterowania/export-cyw-macierz.py','gra/src/game/civ-matrix.ts','gra/src/game/civ-ai-data.ts','gra/src/game/civ-ai-allocation.ts','gra/src/game/civ-bonuses.ts','gra/src/game/economy.ts','gra/src/game/manpower.ts','gra/src/game/population-growth-v85.ts','gra/src/game/diplomacy.ts','gra/src/game/diplomacy-display.ts','gra/src/game/ai.ts','gra/src/game/ai-production-priorities.ts','gra/src/main.ts','docs/decyzje/D-cyw-roster-6-REZERWA.md','Civ-CYWILIZACJE/draft/roster-6-REZERWA.json','Civ-CYWILIZACJE/DOKUMENTACJA-DEV-CYWILIZACJE.md','docs/PORADNIK-GRACZA/13-cywilizacje.md']+list(legacy_docs.values())
def file_info(rel):
    p=ROOT/rel
    if not p.is_file(): return {'path':rel,'exists':False,'bytes':None,'sha256':None}
    b=p.read_bytes(); return {'path':rel,'exists':True,'bytes':len(b),'sha256':sha256(b).hexdigest()}
sources=[file_info(x) for x in source_paths]

def exact_hits(pid):
    result=[]
    for p in sorted((ROOT/'gra'/'src').rglob('*')):
        if p.suffix not in {'.ts','.tsx','.js','.cjs'} or not p.is_file(): continue
        try: lines=p.read_text(encoding='utf-8').splitlines()
        except UnicodeDecodeError: continue
        rel=p.relative_to(ROOT).as_posix()
        result += [f'{rel}:{i}' for i,line in enumerate(lines,1) if pid in line]
    return result

# Workbook helpers.
def fmt(v): return clean(v)
def write_sheet(wb,title,headers,rows,fill_map=None):
    ws=wb.create_sheet(title); ws.sheet_view.showGridLines=False; ws.freeze_panes='A2'; ws.append(headers)
    for c in ws[1]: c.fill=PatternFill('solid',fgColor='1F4E78'); c.font=Font(bold=True,color='FFFFFF'); c.alignment=Alignment(wrap_text=True,vertical='top')
    for row in rows:
        ws.append([fmt(v) for v in row])
        for c in ws[ws.max_row]: c.alignment=Alignment(wrap_text=True,vertical='top'); c.font=Font(size=9)
        if fill_map:
            for c in ws[ws.max_row]:
                if c.value in fill_map:
                    for x in ws[ws.max_row]: x.fill=fill_map[c.value]
                    break
    ws.auto_filter.ref=ws.dimensions
    for i,col in enumerate(ws.iter_cols(min_col=1,max_col=ws.max_column),1):
        n=max((len(str(c.value)) if c.value is not None else 0) for c in list(col)[:80]); ws.column_dimensions[get_column_letter(i)].width=min(52,max(10,n+2))
    ws.row_dimensions[1].height=34
    return ws

wb=Workbook(); wb.remove(wb.worksheets[0]); wb.properties.creator='Hermes Agent'; wb.properties.title='Cyw-macierz — 14 pozostałych cywilizacji'
fill_map={'REAL_GAMEPLAY':PatternFill('solid',fgColor='E2F0D9'),'UI_ONLY':PatternFill('solid',fgColor='FFF2CC'),'DEAD_UNWIRED':PatternFill('solid',fgColor='FCE4D6'),'UNWIRED':PatternFill('solid',fgColor='F4CCCC'),'BLOCKED':PatternFill('solid',fgColor='FCE4D6'),'DECISION_REQUIRED':PatternFill('solid',fgColor='FCE4D6')}
ws=wb.create_sheet('README'); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=30; ws.column_dimensions['B'].width=110
readme=[('Artifact','Civ-macierz-14-pozostale.xlsx'),('Task',TASK),('Generated at',RUN_AT),('Scope','14 target civilizations excluding Grecy; 113 definitions; 1582 cells.'),('Additivity','New workbook; Greece pilot is preserved only in Alokacje_Grecja; no Greece row in Pozostale_14.'),('Primary current source',f"gra/data/civ-matrix.json — _meta.source: {matrix.get('_meta',{}).get('source')}"),('Value rule','Direct target-row lookup. copied_from_greece=false is recorded per cell; equal numbers are not evidence of copying.'),('Provenance rule','CURRENT_RUNTIME_DATA is observed JSON. Historical citations are separate; no ACCEPTED_SPEC claim is invented.'),('Consumer rule','REAL_GAMEPLAY has a traced call-site; UI_ONLY is display-only; DEAD_UNWIRED has non-default data but no gameplay call-site; UNWIRED has no non-default target data and no call-site.'),('Difficulty rule','AI skala fields: Normal±1 clamped 1..10. Growth: Easy×0.5 / Normal×1 / Hard×1.5. Neutral real consumers remain explicit decisions.'),('Review gate','Operator evidence package only; no integration/deploy.'),('Counts',f'consumer cells={dict(status_counts)}; non-default cells={nondefault}; bonus rows={len(bonus_rows)}')]
for i,(a,b) in enumerate(readme,1): ws.cell(i,1,a); ws.cell(i,2,b); ws.cell(i,1).font=Font(bold=True,color='1F4E78'); ws.cell(i,1).fill=PatternFill('solid',fgColor='D9EAF7'); ws.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[i].height=32 if len(str(b))>150 else 22
write_sheet(wb,'Pozostale_14',['Civ ID','Cywilizacja','Tier','Param ID','Domena','Jednostka','Moduł deklarowany','Formula','Wartość bieżąca','Default','Odchylenie','Status wartości','Relacja do defaultu','Provenance','Źródło wartości','Źródło historyczne','Pewność historyczna','Status konsumenta','Easy effective','Normal effective','Hard effective','Polityka poziomów','Formuła efektywna','Kierunek','Live call-site','Uwagi'],matrix_rows,fill_map)
write_sheet(wb,'Dodatkowe_profile',['Civ ID','Cywilizacja','Warstwa','Pole/Param ID','Wartość bieżąca','Wartość adaptera/legacy','Jednostka','Status konsumenta','Easy','Normal','Hard','Polityka','Źródło bieżące','Źródło adaptera','Call-site','Uwagi'],profile_rows,fill_map)
write_sheet(wb,'Aktywne_bonusy',['Civ ID','Cywilizacja','Warstwa','Klucz','Typ/Cel','Wartość','Status konsumenta','Źródło','Live call-site','Polityka poziomów','Opis / uwaga'],bonus_rows,fill_map)
write_sheet(wb,'Konsumenci_i_braki',['Param ID','Domena','Status konsumenta','Polityka poziomów','Przykładowa wartość','Easy','Normal','Hard','Formuła efektywna','Kierunek','Live call-site','Uwagi'],policy_rows,fill_map)
level_rows=[]
for pid in param_ids:
    m=meta(pid); raw=target_rows[0]['params'][pid]; e,n,h=effective(pid,raw,m); gap='DECISION_REQUIRED' if m['policy'] in {'CURRENT_NEUTRAL_MUL_ABS','CURRENT_NEUTRAL_CALLER_OMITS_DIFFICULTY'} else ''
    level_rows.append([pid,m['status'],m['policy'],e,n,h,m['formula'],m['direction'],gap])
write_sheet(wb,'Polityka_poziomow',['Param ID','Status','Policy contract','Current Easy','Current Normal','Current Hard','Formula / rule','Direction','Decision gap'],level_rows,fill_map)
write_sheet(wb,'Alokacje_Grecja',['Civ/pool','Owner kind','Difficulty','Work buildings %','Work empire pool %','Science %','Money %','Wealth %','Automation %','Source','Note'],allocation_rows)
write_sheet(wb,'Historia_zrodel',['Civ ID','Cywilizacja','Tier','Historical status','Citation','Citation note','Confidence','Current matrix row','Current value class','Decision note'],history_rows)
write_sheet(wb,'Nierozstrzygniete',['ID','Obszar','Status','Evidence / gap','Next decision/action','Source'],unresolved,fill_map)
manifest_rows=[[*x.values(),'source evidence'] for x in sources]+[[x,True,'','','generated artifact'] for x in ['panele-sterowania/cyw-macierz/Civ-macierz-14-pozostale.xlsx','panele-sterowania/cyw-macierz/civ-matrix-14-coverage.tsv','panele-sterowania/cyw-macierz/civ-matrix-14-audit.json','docs/decyzje/D-cyw-macierz-14-pozostale.md']]
write_sheet(wb,'Manifest',['Path','Exists','Bytes','SHA-256','Role'],manifest_rows)
xlsx=OUT/'Civ-macierz-14-pozostale.xlsx'; wb.save(xlsx)

# Flat coverage TSV.
tsv=OUT/'civ-matrix-14-coverage.tsv'; headers=list(coverage[0])
with tsv.open('w',encoding='utf-8',newline='') as f:
    f.write('\t'.join(headers)+'\n')
    for row in coverage:
        vals=[]
        for k in headers:
            v=row[k]
            if isinstance(v,bool): v='true' if v else 'false'
            elif isinstance(v,(dict,list)): v=json.dumps(v,ensure_ascii=False,sort_keys=True)
            elif v is None: v=''
            vals.append(str(v).replace('\t',' ').replace('\n',' ').replace('\r',' '))
        f.write('\t'.join(vals)+'\n')

checks={'target_civilization_count':len(target_rows)==14,'parameter_definition_count':len(param_ids)==113,'coverage_cell_count':len(coverage)==1582,'all_target_rows_have_full_schema':all(set(r['params'])==set(param_ids) for r in target_rows),'no_greece_target_row':'grecy' not in target_ids,'direct_row_provenance_for_all_cells':all(c['value_source'].startswith('gra/data/civ-matrix.json#/cywilizacje') and not c['copied_from_greece'] for c in coverage),'xlsx_is_additive_new_name':xlsx.name!='Cyw-macierz-REVIEW.xlsx','reserve_history_not_claimed_accepted':all(history(c)[0]=='RESERVE_DRAFT_UNRESOLVED' for c in reserve_ids)}
assert all(checks.values()), checks

counts=', '.join(f'{k}={v}' for k,v in sorted(status_counts.items()))
doc=ROOT/'docs'/'decyzje'/'D-cyw-macierz-14-pozostale.md'
doc_text=f'''# D-cyw-macierz-14-pozostale — macierz pozostałych 14 cywilizacji

Status: PASS-WITH-NOTES (pakiet dowodowy Operatora; bez integracji i deployu)
Temat: {TASK}
Data generacji: {RUN_AT}

## Zakres

- Źródło: `gra/data/civ-matrix.json`, `_meta.source = "{matrix.get('_meta',{}).get('source')}"`.
- 14 ID poza `grecy`: `{', '.join(target_ids)}`.
- 113 definicji × 14 = **1582 komórki**.
- Workbook jest addytywny. Nie nadpisuje pilota Grecji ani istniejącego workbooka; Grecja jest wyłącznie opisana w `Alokacje_Grecja`.

## Pochodzenie wartości

Każda komórka jest pobrana bezpośrednio z wiersza celu w `gra/data/civ-matrix.json`. `copied_from_greece=false` jest zapisane dla wszystkich 1582 komórek. Równość liczb nie jest traktowana jako dowód kopiowania.

`CURRENT_RUNTIME_DATA` oznacza obserwowaną wartość bieżącego JSON-a. Nie wpisano propozycji do macierzy i nie nadano statusu `ACCEPTED_SPEC` bez jednoznacznego dowodu. Osiem starszych cywilizacji ma tekst encyklopedyczny; sześć rezerwowych ma tylko dowód draft/rezerwa i pozostaje `RESERVE_DRAFT_UNRESOLVED`.

## Konsumenci

- `REAL_GAMEPLAY`: 11 parametrów — 8 pól AI, `lud_wzrost_proc`, `dip_handlowosc_archetyp`, `dip_nastawienie_bazowe`.
- `UI_ONLY`: 5 pól dyplomatycznych używanych przez tagi w `diplomacy-display.ts`.
- Reszta jest `DEAD_UNWIRED` (nie-defaultowe dane bez gameplay call-site) albo `UNWIRED` (brak nie-defaultowych danych i call-site). `paramDefs.modul` nie jest dowodem użycia.
- Bonusy i mnożniki z `civs.json` są w osobnym arkuszu; nie są udawane jako aktywne parametry macierzy.

Komórki wg statusu: {counts}. Komórki różne od defaultu: {nondefault}.

## Polityka poziomów

- AI `skala_1_10`: Normal = macierz; Easy = `max(1, Normal-1)`; Hard = `min(10, Normal+1)`.
- `lud_wzrost_proc`: Easy ×0.5 / Normal ×1 / Hard ×1.5, zaokrąglenie wkładu do punktów procentowych.
- `ai_profil_obronna`: flaga ownera, identyczna na poziomach.
- `dip_handlowosc_archetyp`: realny konsument, ale obecnie neutralny, bo `mul_abs` nie jest skalowane przez helper.
- `dip_nastawienie_bazowe`: realny konsument, lecz bieżący caller pomija difficulty i korzysta z Normal. Oba przypadki pozostają jawnymi decyzjami.

## Otwarte decyzje

1. `HIST-14-RESERVE`: zaakceptować albo pozostawić jako draft sześć pozycji rezerwowych.
2. `WIRE-113`: rozbić niepodłączone pola na osobne tematy implementacyjne.
3. `DIFF-TRADE` i `DIFF-TRUST`: zatwierdzić neutralność albo kierunek Easy/Normal/Hard przed kodem.
4. `ALLOC-14`: nie kopiować alokacji Grecji bez osobnej decyzji.
5. `SOURCE-XLSX`: `_meta.source` wskazuje 11 arkuszy, lecz current base nie zawiera kanonicznego XLSX; wymagane jest osobne źródło archiwalne.

## Artefakty

- `panele-sterowania/cyw-macierz/Civ-macierz-14-pozostale.xlsx`
- `panele-sterowania/cyw-macierz/civ-matrix-14-coverage.tsv`
- `panele-sterowania/cyw-macierz/civ-matrix-14-audit.json`
- `dyspozycje/autobot/runs/{TASK}/00-dispatch.md`
- `dyspozycje/autobot/runs/{TASK}/01-operator.md`

Brak zmian w `gra/src`, commitów, pushu, merge i deployu.
'''
doc.write_text(doc_text,encoding='utf-8')

def manifest(paths):
    result=[]
    for p in paths:
        b=p.read_bytes(); result.append({'path':p.relative_to(ROOT).as_posix(),'bytes':len(b),'sha256':sha256(b).hexdigest()})
    return result

dispatch=RUN/'00-dispatch.md'; dispatch.write_text(f'''STATUS: PASS\nDOMAIN: GAME\nTEMAT: {TASK}\nGOAL: Additive evidence package for 14 non-Greece civilization matrix rows.\nZMIANY/COMMIT: no product-code changes; generated workbook, TSV, JSON and decision/run artifacts.\nTESTY: direct-row schema checks passed; 14 rows, 113 definitions, 1582 cells.\nBLOKADY: reserve historical acceptance, two difficulty policies, canonical XLSX source and unwired fields remain explicit.\nRUNDY: 1/5\nNASTĘPNY KROK: Evaluator review; no integration/deploy in this stage.\nDEPLOY/PUSH: NIE WYKONANO\n''',encoding='utf-8')
operator=RUN/'01-operator.md'; operator.write_text(f'''STATUS: PASS-WITH-NOTES\nDOMAIN: GAME\nTEMAT: {TASK}\nGOAL: Verify 14 rows, 113 definitions, direct provenance, consumer trace and additive workbook.\nZMIANY/COMMIT: generated panele-sterowania/cyw-macierz/Civ-macierz-14-pozostale.xlsx, civ-matrix-14-coverage.tsv, docs/decyzje/D-cyw-macierz-14-pozostale.md; no gra/src changes.\nTESTY: generation self-check PASS; target_rows=14; paramDefs=113; coverage=1582; Greece excluded; direct-row provenance=1582/1582. External readback/git/typecheck follows.\nBLOKADY: HIST-14-RESERVE, DIFF-TRADE, DIFF-TRUST, SOURCE-XLSX, WIRE-113 documented; nothing silently resolved.\nRUNDY: 1/5\nNASTĘPNY KROK: Evaluator/Final Control review.\nDEPLOY/PUSH: NIE WYKONANO\n''',encoding='utf-8')

artifacts=manifest([xlsx,tsv,doc,dispatch,operator])
audit={'task_id':TASK,'generated_at':RUN_AT,'status':'PASS-WITH-NOTES','domain':'GAME','goal':'Additive historical/runtime matrix evidence for the 14 non-Greece civilizations.','base':{'repository_root':str(ROOT),'current_matrix_source':'gra/data/civ-matrix.json','matrix_meta':matrix.get('_meta',{}),'greece_excluded_id':'grecy','target_ids':target_ids},'counts':{'target_civilizations':14,'parameter_definitions':113,'coverage_cells':1582,'non_default_cells':nondefault,'consumer_cell_counts':dict(status_counts),'active_bonus_rows':len(bonus_rows),'active_bonus_status_counts':dict(bonus_status)},'checks':checks,'source_manifest':sources,'historical_sources':{cid:{'status':history(cid)[0],'citation':history(cid)[1],'confidence':history(cid)[2],'note':history(cid)[3]} for cid in target_ids},'consumer_parameter_summary':[{'param_id':pid,'definition':param_defs[pid],'consumer':meta(pid),'exact_source_hits':exact_hits(pid),'target_non_default_count':sum(r['params'][pid]!=defaults.get(pid) for r in target_rows)} for pid in param_ids],'open_decisions':[dict(zip(['id','area','status','evidence','next_action','source'],r)) for r in unresolved],'artifacts':artifacts,'verification':{'generation_self_checks':'PASS','external_checks':'PENDING'}}
(OUT/'civ-matrix-14-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'status':'PASS-WITH-NOTES','task':TASK,'targets':14,'params':113,'coverage':1582,'non_default_cells':nondefault,'consumer_cell_counts':dict(status_counts),'bonus_rows':len(bonus_rows),'xlsx':str(xlsx.relative_to(ROOT)),'coverage_tsv':str(tsv.relative_to(ROOT)),'audit_json':str((OUT/'civ-matrix-14-audit.json').relative_to(ROOT))},ensure_ascii=False))
